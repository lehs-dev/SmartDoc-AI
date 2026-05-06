import fitz  # PyMuPDF
import docx
from langchain_text_splitters import RecursiveCharacterTextSplitter

import fitz  # PyMuPDF
import docx
from langchain_text_splitters import RecursiveCharacterTextSplitter

def extract_text(file_path: str, filename: str) -> str:
    """Trích xuất text từ PDF (xử lý đa cột) hoặc DOCX."""
    text = ""
    ext = filename.split('.')[-1].lower()

    try:
        if ext == 'pdf':
            doc = fitz.open(file_path)
            for page in doc:
                # Lấy text dưới dạng các khối (blocks)
                blocks = page.get_text("blocks")
                
                # blocks là list các tuple: (x0, y0, x1, y1, "text", block_no, block_type)
                # Chỉ lấy block chứa text (block_type == 0)
                text_blocks = [b for b in blocks if b[6] == 0]
                
                # Sắp xếp các block: Ưu tiên cột (trái sang phải - x0), sau đó từ trên xuống (y0)
                # Mẹo: Chia tọa độ x0 thành các dải (vd: cách nhau 100px) để gom cột
                text_blocks.sort(key=lambda b: (round(b[0] / 100), b[1]))
                
                for b in text_blocks:
                    # Nối text của từng block lại, làm sạch khoảng trắng thừa
                    clean_text = b[4].replace('\n', ' ').strip()
                    if clean_text:
                        text += clean_text + "\n\n"
            doc.close()

        elif ext == 'docx':
            # ... (Giữ nguyên logic DOCX của ông) ...
            doc = docx.Document(file_path)
            for para in doc.paragraphs:
                if para.text.strip():
                    text += para.text + "\n"
            for table in doc.tables:
                for row in table.rows:
                    row_data = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_data:
                        text += " | ".join(row_data) + "\n"
                        
    except Exception as e:
        print(f"Lỗi đọc file {filename}: {e}")
        
    return text


def chunk_text(text: str) -> list[str]:
    """
    Băm văn bản thành các khúc nhỏ.
    Cấu hình Chunk size = 300 để ép FAISS tìm ra đúng chỗ chứa số liệu,
    tránh LLM 7B bị ngợp thông tin (hallucinate).
    """
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=300,
        chunk_overlap=50,
        separators=["\n\n", "\n", ".", "!", "?", " ", ""]
    )
    return splitter.split_text(text)